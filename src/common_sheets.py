"""
Shared sheets used by more than one workbook.

Each workbook in this project is self-contained: it restates the public source
data it needs on its own sheet rather than linking across files. Cross-workbook
references break the moment someone renames or moves a file, and a reviewer
should be able to open any one of these and have it work.
"""

from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter as gcl

import inputs as I
import styles as S

HY = I.HIST_YEARS
HCOL = {y: gcl(2 + i) for i, y in enumerate(HY)}

FIELDS = [
    ("Contributions and grants", "contributions", "PUBLIC", "src.contrib"),
    ("Program service revenue", "program_rev", "PUBLIC", "src.progrev"),
    ("Investment and other income", "invest_inc", "PUBLIC", "src.invinc"),
    ("Less: fundraising event costs and asset-sale losses", "revenue_netting", "DERIVED",
     "src.netting"),
    ("Total revenue", "total_revenue", "PUBLIC", "src.rev"),
    ("Total functional expenses", "total_expense", "PUBLIC", "src.exp"),
    ("Officer compensation", "officer_comp", "PUBLIC", "src.offcomp"),
    ("Other salaries and wages", "other_salaries", "PUBLIC", "src.othersal"),
    ("Payroll taxes", "payroll_tax", "PUBLIC", "src.ptax"),
    ("Total assets", "total_assets", "PUBLIC", "src.assets"),
    ("Total liabilities", "total_liab", "PUBLIC", "src.liab"),
]


def source_data_sheet(wb, R: dict, title="Source data"):
    """Public Form 990 figures plus the modelled splits every workbook needs."""
    ws = wb.create_sheet(title)
    S.set_widths(ws, {"A": 46, **{HCOL[y]: 15 for y in HY}, "H": 11, "I": 54})
    ws.sheet_view.showGridLines = False

    r = S.title_block(ws, "Source data",
                      f"{I.ORG_NAME}, EIN {I.ORG_EIN}. Public Form 990 figures and the modelled "
                      "splits this workbook depends on. Restated here so the file stands alone.",
                      width=9)
    r = S.header_row(ws, r, ["Line"] + [f"FY{y}" for y in HY] + ["Tier", "Note"])

    r = S.section(ws, r, "Form 990 as filed", width=9)
    for name, field, tier, key in FIELDS:
        S.label(ws, r, 1, name, size=9.5, indent=1,
                bold=field in ("total_revenue", "total_expense"))
        for y in HY:
            S.put(ws, r, 2 + HY.index(y), I.F990[y][field], fmt=S.MONEY, is_input=True,
                  bold=field in ("total_revenue", "total_expense"))
        S.tier_cell(ws, r, 8, tier)
        R[key] = r
        r += 1

    S.label(ws, r, 1, "Net assets", bold=True, size=10, indent=1)
    for y in HY:
        col = HCOL[y]
        c = ws.cell(row=r, column=2 + HY.index(y),
                    value=f"={col}{R['src.assets']}-{col}{R['src.liab']}")
        c.number_format = S.MONEY
        c.font = Font(name="Calibri", size=10, bold=True)
        c.alignment = Alignment(horizontal="right")
    S.tier_cell(ws, r, 8, "DERIVED")
    R["src.na"] = r
    r += 2

    r = S.section(ws, r, "Modelled functional expense split", width=9)
    for label_, idx, key in (("Program services share", 0, "src.fs_prog"),
                             ("Management and general share", 1, "src.fs_mg"),
                             ("Fundraising share", 2, "src.fs_fr")):
        S.label(ws, r, 1, label_, size=9.5, indent=1)
        for y in HY:
            S.put(ws, r, 2 + HY.index(y), I.FUNCTIONAL_SPLIT[y][idx], fmt=S.PCT1, is_input=True)
        S.tier_cell(ws, r, 8, "SYNTHETIC")
        R[key] = r
        r += 1
    r += 1

    r = S.section(ws, r, "Derived expense lines", width=9)
    for label_, share_key, key in (("Program services", "src.fs_prog", "src.prog"),
                                   ("Management and general", "src.fs_mg", "src.mg"),
                                   ("Fundraising", "src.fs_fr", "src.fr")):
        S.label(ws, r, 1, label_, size=9.5, indent=1)
        for y in HY:
            col = HCOL[y]
            c = ws.cell(row=r, column=2 + HY.index(y),
                        value=f"={col}{R['src.exp']}*{col}{R[share_key]}")
            c.number_format = S.MONEY
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(horizontal="right")
        R[key] = r
        r += 1
    r += 1

    # Anchors are laid out before the series that reference them, so the row
    # numbers in the interpolation formulas are known rather than predicted.
    p22, p25 = I.IMPACT_PUBLIC[2022]["participants"], I.IMPACT_PUBLIC[2025]["participants"]
    l22, l25 = I.IMPACT_PUBLIC[2022]["placements"], I.IMPACT_PUBLIC[2025]["placements"]

    r = S.section(ws, r, "Interpolation anchors (published observations)", width=9)
    anchor_rows = {}
    for name, val, fmt, key in (("2022 participants", p22, S.NUM, "a.p22"),
                                ("2025 participants", p25, S.NUM, "a.p25"),
                                ("Participant growth per year", None, S.PCT1, "a.pg"),
                                ("2022 placements", l22, S.NUM, "a.l22"),
                                ("2025 placements", l25, S.NUM, "a.l25"),
                                ("Placement growth per year", None, S.PCT1, "a.lg")):
        S.label(ws, r, 1, name, size=9.5, indent=1)
        if val is None:
            S.put(ws, r, 2, f"=(B{r - 1}/B{r - 2})^(1/3)-1", fmt=fmt)
        else:
            S.put(ws, r, 2, val, fmt=fmt, is_input=True)
        S.tier_cell(ws, r, 8, "PUBLIC" if val is not None else "DERIVED")
        anchor_rows[key] = r
        r += 1
    r += 1

    r = S.section(ws, r, "Program volume", width=9)
    for label_, base_key, growth_key, key in (
            ("Participants served", "a.p22", "a.pg", "src.parts"),
            ("Successful outcomes", "a.l22", "a.lg", "src.places")):
        S.label(ws, r, 1, label_, size=9.5, indent=1)
        br, gr = anchor_rows[base_key], anchor_rows[growth_key]
        for y in HY:
            S.put(ws, r, 2 + HY.index(y), f"=$B${br}*(1+$B${gr})^({y}-2022)", fmt=S.NUM1)
        S.tier_cell(ws, r, 8, "DERIVED")
        R[key] = r
        r += 1
    r += 1
    S.note(ws, r, "Participant and placement counts are published for 2022 and 2025 only. Other "
                  "years are interpolated at constant growth between those two observations, which "
                  "fixes the series with no free parameters. The back-cast to FY2019 through "
                  "FY2021 is the weakest link and should be read as illustrative.", width=9)

    S.finish(ws, S.MUTED)
    return ws, r


def disclaimer(ws, row, width=8):
    c = ws.cell(row=row, column=1, value=(
        "Independent analysis prepared from public information for portfolio and educational "
        f"purposes. Not affiliated with, endorsed by, or reviewed by {I.ORG_NAME}. Historical "
        "totals are as filed; all forward-looking and program-level figures are analyst "
        "assumptions and are not the organization's plans or guidance."))
    c.font = Font(name="Calibri", size=8.5, italic=True, color=S.WARN)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    ws.row_dimensions[row].height = 30
    return row + 1
