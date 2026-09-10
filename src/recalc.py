"""
Recalculate a workbook with LibreOffice and report any formula errors.

openpyxl writes formulas but never evaluates them, so a workbook can look
perfect and be full of #REF!. This converts the file headlessly, which forces a
full recalculation, then reads the cached values back and reports every error
cell it finds.
"""

import subprocess
import sys
import shutil
from pathlib import Path

from openpyxl import load_workbook

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "Err:")
SCRATCH = Path("/tmp/claude-0/-home-claude/06d154c5-108f-5eea-9706-a1170a82eaaa/scratchpad/recalc")


def recalc(path: Path) -> Path:
    """Return a path to a recalculated copy carrying cached values."""
    SCRATCH.mkdir(parents=True, exist_ok=True)
    out = SCRATCH / path.name
    if out.exists():
        out.unlink()
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx",
         "--outdir", str(SCRATCH), str(path)],
        check=True, capture_output=True, timeout=300,
    )
    if not out.exists():
        raise RuntimeError(f"LibreOffice produced no output for {path}")
    return out


def scan(path: Path, verbose: bool = True) -> dict:
    """Recalculate and report error cells plus a value map."""
    calc = recalc(path)
    wb = load_workbook(calc, data_only=True)
    problems = []
    values = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                values[f"{ws.title}!{cell.coordinate}"] = v
                if isinstance(v, str) and any(v.startswith(e) or v == e for e in ERRORS):
                    problems.append((ws.title, cell.coordinate, v))
    if verbose:
        if problems:
            print(f"  {path.name}: {len(problems)} error cell(s)")
            for sheet, coord, v in problems[:40]:
                print(f"    {sheet}!{coord} -> {v}")
            if len(problems) > 40:
                print(f"    ... and {len(problems) - 40} more")
        else:
            print(f"  {path.name}: clean, {len(values)} populated cells")
    return dict(path=calc, problems=problems, values=values, workbook=wb)


def get(res: dict, sheet: str, coord: str):
    return res["values"].get(f"{sheet}!{coord}")


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    targets = sys.argv[1:] or sorted(str(p) for p in root.glob("excel-models/*.xlsx"))
    total = 0
    for t in targets:
        res = scan(Path(t))
        total += len(res["problems"])
    print(f"\ntotal error cells: {total}")
    sys.exit(1 if total else 0)
