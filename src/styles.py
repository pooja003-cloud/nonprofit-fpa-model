"""Shared openpyxl formatting so every workbook looks like it came from the same finance team."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Palette - restrained, print-safe, readable on a projector
INK = "1A1A1A"
MUTED = "5F6B7A"
NAVY = "12324F"
NAVY_LIGHT = "E8EEF4"
RULE = "C7D2DC"
ACCENT = "0B6E4F"
WARN = "9A3412"
INPUT_BLUE = "0033CC"     # finance convention: blue text = hardcoded input
BAND = "F5F8FA"

THIN = Side(style="thin", color=RULE)
MEDIUM = Side(style="medium", color=NAVY)

MONEY = '#,##0;[Red](#,##0);"-"'
MONEY0 = '$#,##0;[Red]($#,##0);"-"'
MONEY2 = '#,##0.00;[Red](#,##0.00);"-"'
PCT1 = '0.0%;[Red](0.0%);"-"'
PCT2 = '0.00%;[Red](0.00%);"-"'
NUM = '#,##0;[Red](#,##0);"-"'
NUM1 = '#,##0.0;[Red](#,##0.0);"-"'
MULT = '0.00"x"'
MONTHS = '0.0" mo"'


def title_block(ws: Worksheet, title: str, subtitle: str = "", row: int = 1, width: int = 10):
    """Standard workbook header."""
    c = ws.cell(row=row, column=1, value=title)
    c.font = Font(name="Calibri", size=15, bold=True, color=NAVY)
    ws.cell(row=row, column=1).alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24
    if subtitle:
        s = ws.cell(row=row + 1, column=1, value=subtitle)
        s.font = Font(name="Calibri", size=9.5, italic=True, color=MUTED)
        s.alignment = Alignment(vertical="top", wrap_text=False)
    for col in range(1, width + 1):
        ws.cell(row=row + 2, column=col).border = Border(bottom=MEDIUM)
    return row + 3


def section(ws: Worksheet, row: int, text: str, width: int = 10):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col in range(1, width + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        if col > 1 and cell.value is None:
            cell.value = None
    ws.row_dimensions[row].height = 18
    return row + 1


def header_row(ws: Worksheet, row: int, labels: list, start_col: int = 1, widths: dict | None = None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        c.fill = PatternFill("solid", fgColor=NAVY_LIGHT)
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)
        c.border = Border(top=THIN, bottom=Side(style="medium", color=NAVY), left=THIN, right=THIN)
    ws.row_dimensions[row].height = 30
    return row + 1


def label(ws, row, col, text, bold=False, indent=0, italic=False, color=INK, size=10):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=color)
    c.alignment = Alignment(horizontal="left", indent=indent, vertical="center")
    return c


def put(ws, row, col, value, fmt=MONEY, bold=False, is_input=False, italic=False, color=None):
    """Write a value or formula. Blue text marks a hardcoded input, per finance convention."""
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = fmt
    c.font = Font(name="Calibri", size=10, bold=bold, italic=italic,
                  color=color or (INPUT_BLUE if is_input else INK))
    c.alignment = Alignment(horizontal="right", vertical="center")
    return c


def total_row_style(ws, row, cols, double=False):
    for col in cols:
        c = ws.cell(row=row, column=col)
        c.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        c.border = Border(top=Side(style="thin", color=NAVY),
                          bottom=Side(style="double", color=NAVY) if double else None)


def band(ws, row, cols):
    for col in cols:
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BAND)


def note(ws, row, text, col=1, width=10):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name="Calibri", size=9, italic=True, color=MUTED)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=width)
    ws.row_dimensions[row].height = max(14, 13 * (len(text) // (width * 13) + 1))
    return row + 1


def set_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col if isinstance(col, str) else get_column_letter(col)].width = w


def freeze(ws, cell: str):
    ws.freeze_panes = cell


def finish(ws, tab_color: str = NAVY, gridlines: bool = False):
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = gridlines


def provenance_legend(ws, row, width=10):
    """Small legend explaining the provenance colour coding."""
    from provenance import TIER_ORDER, TIER_COLOR, TIER_DESCRIPTION
    label(ws, row, 1, "Provenance key", bold=True, size=10)
    row += 1
    for t in TIER_ORDER:
        c = ws.cell(row=row, column=1, value=t)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TIER_COLOR[t])
        c.alignment = Alignment(horizontal="center")
        d = ws.cell(row=row, column=2, value=TIER_DESCRIPTION[t])
        d.font = Font(name="Calibri", size=9, italic=True, color=MUTED)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=width)
        row += 1
    return row + 1


def tier_cell(ws, row, col, tier):
    from provenance import TIER_COLOR
    c = ws.cell(row=row, column=col, value=tier)
    c.font = Font(name="Calibri", size=8.5, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TIER_COLOR[tier])
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c
