"""
Model integrity tests.

Two independent implementations of the same model - live Excel formulas and the
Python engine - and these tests check that they agree. That is a much stronger
claim than either being internally consistent.

Every workbook is recalculated headlessly with LibreOffice, which forces Excel to
evaluate the formulas openpyxl wrote but never computed. Without that step a
workbook full of #REF! errors would pass every test.

    python -m pytest tests -q

The recalculation is slow (a few seconds per workbook), so it is cached at
session scope.
"""

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import inputs as I           # noqa: E402
import engine                # noqa: E402
import recalc                # noqa: E402

TOL = 1e-6


def rel(a, b):
    return abs(a - b) / max(abs(b), 1e-9)


# ==========================================================================
# Fixtures - build then recalculate each workbook once per session
# ==========================================================================

@pytest.fixture(scope="session")
def wb03():
    import build_03_financial_model as B
    B.build()
    return B, recalc.scan(ROOT / "excel-models" / "02_Financial_Model.xlsx", verbose=False)


@pytest.fixture(scope="session")
def wb04():
    import build_04_budget_vs_actual as B
    B.build()
    return B, recalc.scan(ROOT / "excel-models" / "03_Budget_vs_Actual.xlsx", verbose=False)


@pytest.fixture(scope="session")
def wb05():
    import build_05_scenario_model as B
    B.build()
    return B, recalc.scan(ROOT / "excel-models" / "04_Scenario_Model.xlsx", verbose=False)


@pytest.fixture(scope="session")
def wb06():
    import build_06_program_economics as B
    B.build()
    return B, recalc.scan(ROOT / "excel-models" / "05_Program_Economics.xlsx", verbose=False)


@pytest.fixture(scope="session")
def wb07():
    import build_07_resource_allocation as B
    B.build()
    return B, recalc.scan(ROOT / "excel-models" / "06_Resource_Allocation.xlsx", verbose=False)


@pytest.fixture(scope="session")
def wb02():
    import build_02_assumptions as B
    B.build()
    return B, recalc.scan(ROOT / "excel-models" / "01_Assumptions.xlsx", verbose=False)


ALL_WB = ["wb02", "wb03", "wb04", "wb05", "wb06", "wb07"]


# ==========================================================================
# 1. No workbook contains a formula error
# ==========================================================================

@pytest.mark.parametrize("fixture", ALL_WB)
def test_no_formula_errors(fixture, request):
    _, res = request.getfixturevalue(fixture)
    assert res["problems"] == [], (
        f"{len(res['problems'])} error cell(s): "
        + ", ".join(f"{s}!{c}={v}" for s, c, v in res['problems'][:10]))


# ==========================================================================
# 2. The public data is what it claims to be
# ==========================================================================

def test_990_revenue_components_reconcile():
    """Gross components plus the netting line must equal filed total revenue."""
    for y in I.HIST_YEARS:
        d = I.F990[y]
        total = (d["contributions"] + d["program_rev"] + d["invest_inc"]
                 + d["revenue_netting"])
        assert abs(total - d["total_revenue"]) < 1, f"FY{y} revenue does not tie"


def test_fy2024_net_assets_tie_to_published():
    d = I.F990[2024]
    assert abs((d["total_assets"] - d["total_liab"]) - I.NET_ASSETS_2024_PUBLIC) < 1


def test_functional_split_sums_to_one():
    for y, (p, m, f) in I.FUNCTIONAL_SPLIT.items():
        assert abs(p + m + f - 1.0) < 1e-9, f"FY{y} functional split does not sum to 100%"


def test_revenue_mix_sums_to_one():
    assert abs(sum(I.REVENUE_MIX.values()) - 1.0) < 1e-9


def test_balance_sheet_sums_to_filed_assets():
    assert abs(sum(I.BALANCE_SHEET_2024.values()) - I.F990[2024]["total_assets"]) < 1


def test_every_input_has_a_source():
    """The whole discipline of the project rests on this."""
    for inp in I.REGISTER:
        assert inp.source.strip(), f"{inp.key} has no source"
        assert inp.tier in ("PUBLIC", "DERIVED", "BENCHMARK", "SYNTHETIC")


def test_synthetic_inputs_are_disclosed_not_hidden():
    counts = I.REGISTER.tier_counts()
    assert counts["SYNTHETIC"] > 0, "a model with no assumptions is not being honest"
    assert counts["PUBLIC"] > 0, "a model with no public data is not grounded"


# ==========================================================================
# 3. Program portfolio reconciles to the public control totals
# ==========================================================================

def test_program_participants_tie_to_derived_total():
    pe = engine.program_economics()
    assert abs(pe["totals"]["participants_vs_derived"]) < 1


def test_program_outcomes_tie_to_derived_total():
    pe = engine.program_economics()
    assert abs(pe["totals"]["outcomes_vs_derived"]) < 1


def test_program_costs_tie_to_expense_pool():
    pe = engine.program_economics()
    assert pe["totals"]["program_pool_reconciled"]


def test_program_hours_tie_to_staffing_model():
    hours = sum(I.PROGRAM_BASE[p]["participants"]
                * I.PROGRAM_BASE[p]["staff_hours_per_participant"] for p in I.PROGRAMS)
    target = I.FTE_2024 * I.PROGRAM_FTE_SHARE * I.STAFF_HOURS_PER_FTE
    assert abs(hours - target) < 1


def test_blended_salary_gain_matches_published_within_one_percent():
    """The one external check available on the synthetic program mix."""
    t = engine.program_economics()["totals"]
    assert abs(t["calibration_variance"]) < 0.01, (
        f"blended gain {t['blended_salary_gain']:,.0f} vs published "
        f"{t['published_salary_gain']:,.0f}")


# ==========================================================================
# 4. Excel agrees with Python - historical
# ==========================================================================

HIST_PAIRS = [
    ("h.rev", "revenue"), ("h.exp", "expenses"), ("h.result", "operating_result"),
    ("h.contrib", "contributions"), ("h.prog", "program_expense"), ("h.mg", "mg_expense"),
    ("h.fr", "fundraising_expense"), ("h.na", "net_assets"), ("h.cash", "cash"),
    ("h.liquid", "liquid_reserves"), ("h.personnel", "personnel"),
    ("h.progratio", "program_ratio"), ("h.adminratio", "admin_ratio"),
    ("h.frratio", "fundraising_cost_ratio"), ("h.frroi", "fundraising_roi"),
    ("h.cashrunway", "cash_runway_months"), ("h.liqrunway", "liquid_runway_months"),
    ("h.namult", "net_asset_multiple"), ("h.parts", "participants"),
    ("h.places", "placements"), ("h.outrate", "placement_rate"),
    ("h.cpp", "cost_per_participant"), ("h.cpo", "cost_per_outcome"),
    ("h.margin", "operating_margin"), ("h.persratio", "personnel_ratio"),
]


@pytest.mark.parametrize("key,field", HIST_PAIRS)
def test_excel_historical_matches_engine(wb03, key, field):
    B, res = wb03
    hist = engine.historical()["rows"]
    for y in I.HIST_YEARS:
        xl = recalc.get(res, "Historical analysis", B.HCOL[y] + str(B.R[key]))
        py = hist[y][field]
        assert xl is not None, f"{field} FY{y} is empty in Excel"
        assert rel(xl, py) < TOL, f"{field} FY{y}: excel {xl} vs python {py}"


def test_excel_historical_revenue_ties_to_filing(wb03):
    B, res = wb03
    for y in I.HIST_YEARS:
        xl = recalc.get(res, "Historical analysis", B.HCOL[y] + str(B.R["h.rev"]))
        assert abs(xl - I.F990[y]["total_revenue"]) < 1


# ==========================================================================
# 5. Excel agrees with Python - forecast, all three scenarios
# ==========================================================================

FCST_PAIRS = [
    ("f.rev", "total_revenue"), ("f.exp", "total_expense"), ("f.result", "operating_result"),
    ("f.contrib", "contributions"), ("f.progrev", "program_revenue"),
    ("f.invinc", "investment_income"), ("f.prog", "program_expense"),
    ("f.mg", "mg_expense"), ("f.fr", "fundraising_expense"),
    ("f.liq_close", "closing_liquid"), ("f.na", "closing_net_assets"),
    ("f.parts", "participants"), ("f.out", "outcomes"), ("f.fte", "fte"),
    ("f.personnel", "personnel"),
]


@pytest.mark.parametrize("key,field", FCST_PAIRS)
def test_excel_forecast_matches_engine(wb03, key, field):
    B, res = wb03
    fc = {r.year: r for r in engine.forecast("Base")}
    for y in I.FCST_YEARS:
        xl = recalc.get(res, "Forecast", B.FCOL[y] + str(B.R[key]))
        py = getattr(fc[y], field)
        assert xl is not None, f"{field} FY{y} is empty in Excel"
        assert rel(xl, py) < TOL, f"{field} FY{y}: excel {xl} vs python {py}"


SCEN_PAIRS = [
    ("s.rev", "total_revenue"), ("s.exp", "total_expense"), ("s.result", "operating_result"),
    ("s.contrib", "contributions"), ("s.prog", "program_expense"), ("s.mg", "mg_expense"),
    ("s.fr", "fundraising_expense"), ("s.liq_close", "closing_liquid"),
    ("s.na", "closing_net_assets"), ("s.parts", "participants"), ("s.out", "outcomes"),
]


@pytest.mark.parametrize("key,field", SCEN_PAIRS)
def test_scenario_workbook_base_case_matches_engine(wb05, key, field):
    B, res = wb05
    fc = {r.year: r for r in engine.forecast("Base", years=B.YEARS)}
    for y in B.YEARS:
        xl = recalc.get(res, "Scenario model", B.YCOL[y] + str(B.R[key]))
        py = getattr(fc[y], field)
        assert xl is not None
        assert rel(xl, py) < TOL, f"{field} FY{y}: excel {xl} vs python {py}"


@pytest.mark.parametrize("scenario", ["Downside", "Upside"])
def test_scenario_switch_drives_the_model(wb05, scenario, tmp_path):
    """
    The CHOOSE-driven scenario switch is the headline feature of 05. This flips
    it and confirms the entire five-year model re-runs to match the engine.
    """
    from openpyxl import load_workbook
    B, _ = wb05
    wb = load_workbook(ROOT / "excel-models" / "04_Scenario_Model.xlsx")
    wb["Scenario switch"].cell(row=B.R["sw.sel"], column=2).value = scenario
    target = tmp_path / f"switch_{scenario}.xlsx"
    wb.save(target)
    res = recalc.scan(target, verbose=False)
    assert res["problems"] == []
    fc = {r.year: r for r in engine.forecast(scenario, years=B.YEARS)}
    for key, field in SCEN_PAIRS:
        for y in B.YEARS:
            xl = recalc.get(res, "Scenario model", B.YCOL[y] + str(B.R[key]))
            py = getattr(fc[y], field)
            assert rel(xl, py) < TOL, f"{scenario} {field} FY{y}: excel {xl} vs python {py}"


# ==========================================================================
# 6. Budget vs actual and the variance bridge
# ==========================================================================

BVA_PAIRS = [
    ("b.rev", "Total revenue"), ("b.exp", "Total expenses"),
    ("b.prog", "  Program services"), ("b.mg", "  Management and general"),
    ("b.fr", "  Fundraising"), ("b.contrib", "  Contributions and grants"),
    ("b.result", "Operating result"), ("b.parts", "Participants served"),
    ("b.places", "Successful outcomes"), ("b.cpp", "Cost per participant"),
    ("b.personnel", "Personnel (memo)"),
]


@pytest.mark.parametrize("key,label", BVA_PAIRS)
def test_budget_vs_actual_matches_engine(wb04, key, label):
    B, res = wb04
    L = engine.budget_vs_actual()["lines"]
    for col, field in (("B", "budget"), ("C", "actual")):
        xl = recalc.get(res, "Budget vs actual", col + str(B.R[key]))
        py = L[label][field]
        assert xl is not None
        assert rel(xl, py) < TOL, f"{label} {field}: excel {xl} vs python {py}"


def test_actual_column_ties_to_the_filing(wb04):
    """The actual column is real data and must equal the 990 exactly."""
    B, res = wb04
    rev = recalc.get(res, "Budget vs actual", "C" + str(B.R["b.rev"]))
    exp = recalc.get(res, "Budget vs actual", "C" + str(B.R["b.exp"]))
    assert abs(rev - I.F990[2024]["total_revenue"]) < 1
    assert abs(exp - I.F990[2024]["total_expense"]) < 1


def test_variance_bridge_reconciles(wb04):
    B, res = wb04
    tot = recalc.get(res, "Variance decomposition", "B" + str(B.R["d.totvar"]))
    vol = recalc.get(res, "Variance decomposition", "B" + str(B.R["d.vol"]))
    rate = recalc.get(res, "Variance decomposition", "B" + str(B.R["d.rate"]))
    assert abs(vol + rate - tot) < 0.01, "volume plus rate must equal the total variance"
    assert recalc.get(res, "Variance decomposition", "B" + str(B.R["d.check"])) == "ok"


def test_variance_decomposition_matches_engine(wb04):
    B, res = wb04
    dec = engine.budget_vs_actual()["decomposition"]
    for key, field in (("d.totvar", "total_program_variance"), ("d.vol", "volume_variance"),
                       ("d.rate", "rate_variance")):
        xl = recalc.get(res, "Variance decomposition", "B" + str(B.R[key]))
        assert abs(xl - dec[field]) < 0.01


# ==========================================================================
# 7. Program economics
# ==========================================================================

PROG_PAIRS = [
    ("e.parts", "participants"), ("e.comp", "completers"), ("e.out", "outcomes"),
    ("e.cost", "program_cost"), ("e.cpp", "cost_per_participant"),
    ("e.cpc", "cost_per_completer"), ("e.cpo", "cost_per_outcome"),
    ("e.outrate", "outcome_rate"), ("e.direct", "direct_cost"),
    ("e.indirect", "indirect_cost"), ("e.hours", "staff_hours"),
    ("e.util", "capacity_utilization"), ("e.totgain", "total_earnings_gain"),
    ("e.gainperdollar", "earnings_gain_per_dollar"),
]


@pytest.mark.parametrize("key,field", PROG_PAIRS)
def test_program_economics_matches_engine(wb06, key, field):
    B, res = wb06
    pe = engine.program_economics()["programs"]
    for p in I.PROGRAMS:
        xl = recalc.get(res, "Program economics", B.PCOL[p] + str(B.R[key]))
        py = pe[p][field]
        assert xl is not None
        assert rel(xl, py) < TOL, f"{p} {field}: excel {xl} vs python {py}"


def test_program_workbook_reconciliation_checks_pass(wb06):
    B, res = wb06
    ws = res["workbook"]["Program inputs"]
    fails = [c.value for row in ws.iter_rows() for c in row if c.value == "FAIL"]
    assert not fails, "a reconciliation check on 'Program inputs' failed"


# ==========================================================================
# 8. Resource allocation - Excel objective must match the LP solver
# ==========================================================================

def test_allocation_objective_matches_linear_program(wb07):
    B, res = wb07
    opt = engine.allocate(objective="outcomes")
    xl = recalc.get(res, "Allocation model", "B" + str(B.R["a.obj"]))
    assert rel(xl, opt["total_additional_outcomes"]) < 1e-4, (
        f"Excel objective {xl} vs CBC solution {opt['total_additional_outcomes']}")


def test_allocation_is_feasible(wb07):
    B, res = wb07
    flag = recalc.get(res, "Allocation model", "F" + str(B.R["a.con_last"] + 2))
    assert flag == "FEASIBLE"


def test_no_constraint_is_violated(wb07):
    B, res = wb07
    violated = []
    for r in range(B.R["a.con_first"], B.R["a.con_last"] + 1):
        if recalc.get(res, "Allocation model", "F" + str(r)) == "VIOLATED":
            violated.append(recalc.get(res, "Allocation model", "A" + str(r)))
    assert not violated, f"violated constraints: {violated}"


def test_lp_respects_every_constraint():
    """Check the Python solution directly, independent of Excel."""
    a = engine.allocate(objective="outcomes")
    c = a["constraints"]
    pe = engine.program_economics()["programs"]
    total = sum(d["allocation"] for d in a["detail"].values())
    assert total <= a["deployable"] + 0.01
    for p, d in a["detail"].items():
        assert d["allocation"] >= c["min_share_per_program"] * a["deployable"] - 0.01
        assert d["allocation"] <= c["max_share_per_program"] * a["deployable"] + 0.01
        assert d["new_participants"] <= pe[p]["max_participants"] + 0.01
    assert a["total_additional_hours"] <= c["staff_hours_available"] + 0.01


def test_optimised_allocation_beats_naive_policies():
    a = engine.allocate(objective="outcomes")
    assert a["comparison"]["optimized"] >= a["comparison"]["equal_split"]
    assert a["comparison"]["optimized"] >= a["comparison"]["pro_rata"]


def test_objectives_genuinely_disagree():
    """If they agreed, presenting three of them would be theatre."""
    comp = engine.allocation_objective_comparison()
    assert comp["outcomes_price_of_earnings_focus"] > 1
    assert comp["earnings_price_of_outcome_focus"] > 1


def test_balanced_case_sits_between_the_extremes():
    comp = engine.allocation_objective_comparison()
    r = {row["objective"]: row for row in comp["rows"]}
    assert r["earnings"]["outcomes"] < r["balanced"]["outcomes"] <= r["outcomes"]["outcomes"]
    assert r["outcomes"]["avg_gain"] <= r["balanced"]["avg_gain"] < r["earnings"]["avg_gain"]


# ==========================================================================
# 9. Assumptions workbook reconciliation
# ==========================================================================

def test_assumptions_reconciliation_checks_pass(wb02):
    B, res = wb02
    ws = res["workbook"]["Reconciliation checks"]
    fails = [c.coordinate for row in ws.iter_rows() for c in row if c.value == "FAIL"]
    assert not fails, f"reconciliation failures at {fails}"


# ==========================================================================
# 10. Model behaviour - the forecast should behave sensibly
# ==========================================================================

def test_cash_rollforward_closes():
    """Opening plus result must equal closing, every year, every scenario."""
    for scen in I.SCENARIOS:
        for r in engine.forecast(scen, years=engine.SENSITIVITY_HORIZON):
            assert abs(r.opening_liquid + r.operating_result - r.closing_liquid) < 0.01


def test_opening_balance_chains_across_years():
    for scen in I.SCENARIOS:
        rows = engine.forecast(scen, years=engine.SENSITIVITY_HORIZON)
        for prev, cur in zip(rows, rows[1:]):
            assert abs(prev.closing_liquid - cur.opening_liquid) < 0.01


def test_expense_components_sum_to_total():
    for scen in I.SCENARIOS:
        for r in engine.forecast(scen, years=engine.SENSITIVITY_HORIZON):
            total = r.program_expense + r.mg_expense + r.fundraising_expense
            assert abs(total - r.total_expense) < 0.01


def test_revenue_components_sum_to_total():
    for scen in I.SCENARIOS:
        for r in engine.forecast(scen, years=engine.SENSITIVITY_HORIZON):
            total = r.contributions + r.program_revenue + r.investment_income
            assert abs(total - r.total_revenue) < 0.01


def test_program_costs_sum_to_program_expense():
    for scen in I.SCENARIOS:
        for r in engine.forecast(scen, years=engine.SENSITIVITY_HORIZON):
            assert abs(sum(r.cost_by_program.values()) - r.program_expense) < 0.01


def test_scenarios_are_ordered():
    """Upside must beat Base must beat Downside, or the drivers are miswired."""
    last = {s: engine.forecast(s, years=engine.SENSITIVITY_HORIZON)[-1] for s in I.SCENARIOS}
    assert last["Downside"].total_revenue < last["Base"].total_revenue < last["Upside"].total_revenue
    assert last["Downside"].operating_result < last["Upside"].operating_result
    assert last["Downside"].outcomes < last["Base"].outcomes < last["Upside"].outcomes


def test_participants_never_exceed_capacity():
    for scen in I.SCENARIOS:
        for r in engine.forecast(scen, years=engine.SENSITIVITY_HORIZON):
            for p, v in r.participants_by_program.items():
                assert v <= I.PROGRAM_BASE[p]["max_participants"] + 0.01


def test_funding_shock_reduces_revenue_monotonically():
    prev = None
    for shock in (0.0, 0.1, 0.2, 0.3, 0.4):
        f = engine.forecast("Base", funding_shock=shock, shock_year=2025)
        rev = f[0].total_revenue
        if prev is not None:
            assert rev < prev, "a larger funding cut must reduce revenue"
        prev = rev


def test_reserve_floor_enforcement_cuts_programs():
    """A shock large enough to threaten the floor must reduce beneficiaries."""
    years = engine.SENSITIVITY_HORIZON
    ref = engine.forecast("Base", years=years, enforce_reserve_floor=True)
    hit = engine.forecast("Base", funding_shock=0.60, shock_year=years[0], years=years,
                          enforce_reserve_floor=True)
    assert hit[-1].participants < ref[-1].participants
    assert hit[-1].outcomes < ref[-1].outcomes
    assert max(r.contraction_applied for r in hit) > 0


def test_reserve_floor_is_actually_held():
    years = engine.SENSITIVITY_HORIZON
    floor = I.ALLOCATION_CONSTRAINTS["reserve_floor_months"]
    for shock in (0.4, 0.5, 0.6):
        f = engine.forecast("Base", funding_shock=shock, shock_year=years[0], years=years,
                            enforce_reserve_floor=True)
        for r in f:
            runway = r.closing_liquid / (r.total_expense / 12)
            assert runway >= floor - 0.01, f"floor breached at shock {shock}: {runway:.2f} months"


def test_minimum_viable_funding_is_a_real_boundary():
    """Just inside the threshold holds the floor; just outside breaks it."""
    mvf = engine.minimum_viable_funding("Base")
    years = engine.SENSITIVITY_HORIZON
    floor = mvf["reserve_floor_months"]

    def worst(shock):
        f = engine.forecast("Base", funding_shock=shock, shock_year=years[0], years=years)
        return min(r.closing_liquid / (r.total_expense / 12) for r in f)

    assert worst(mvf["max_sustainable_shock"] - 0.005) >= floor - 0.05
    assert worst(mvf["max_sustainable_shock"] + 0.01) < floor


def test_sroi_is_reported_as_a_range_not_a_point():
    rng = engine.sroi_sensitivity()
    assert rng["spread"] > 2, (
        "if the SROI ratio were stable across assumptions, the caveats would be overdone; "
        "this test exists to confirm the instability the workbook warns about is real")


# ==========================================================================
# 11. Provenance discipline
# ==========================================================================

def test_no_forward_looking_figure_is_tagged_public():
    for inp in I.REGISTER:
        if inp.section.startswith(("10.", "11.", "12.", "13.")):
            assert inp.tier != "PUBLIC", (
                f"{inp.key} is forward-looking or an assumption but tagged PUBLIC")


def test_benchmarks_cite_a_named_standard():
    for inp in I.REGISTER:
        if inp.tier == "BENCHMARK":
            assert any(s in inp.source for s in ("give.org", "propelnonprofits", "BBB", "Propel"))


# ==========================================================================
# 12. The Power BI workbook is importable
#
# Power BI rejected an earlier version of this file with "We were unable to
# load this Excel file because we couldn't understand its format". The data was
# fine; the problem was the shape of the XML. openpyxl's fast writer (used
# whenever lxml is installed) emits numeric cells as t="n" and strings as
# inline runs with no shared string table. That is legal OOXML, but it is not
# what Excel writes, and Power Query's parser refuses it - with an error that
# names no sheet, no column and no cell.
#
# These tests pin the file's shape, not just its contents, because the contents
# were never what was wrong.
# ==========================================================================

import csv as _csv           # noqa: E402
import re as _re             # noqa: E402
import math as _math         # noqa: E402
import zipfile               # noqa: E402

PBI_XLSX = ROOT / "powerbi" / "PowerBI_Data_Model.xlsx"
PBI_DATA = ROOT / "powerbi" / "data"


@pytest.fixture(scope="session")
def pbi_built():
    import build_08_powerbi as B
    B.build()
    return B


def test_powerbi_workbook_exists(pbi_built):
    assert PBI_XLSX.exists(), "the browser-import workbook was not generated"


def test_powerbi_workbook_uses_a_shared_string_table(pbi_built):
    """Excel always writes one. Power Query expects one."""
    with zipfile.ZipFile(PBI_XLSX) as z:
        assert "xl/sharedStrings.xml" in z.namelist(), (
            "no shared string table - the workbook was written by openpyxl's fast "
            "path again, and Power BI will reject it")


def test_powerbi_workbook_has_no_inline_strings_or_typed_numbers(pbi_built):
    """The two specific markers of the writer that Power Query refused."""
    with zipfile.ZipFile(PBI_XLSX) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert sheets
        for n in sheets:
            xml = z.read(n).decode()
            assert "inlineStr" not in xml, f"{n} uses inline strings"
            assert 't="n"' not in xml, f'{n} writes t="n" on numeric cells'


def _tables(wb):
    """Map table name -> worksheet. The table name is what Power BI imports."""
    out = {}
    for ws in wb.worksheets:
        for t in ws.tables:
            out[t] = ws
    return out


def test_powerbi_workbook_tables_are_named_excel_tables(pbi_built):
    from openpyxl import load_workbook
    wb = load_workbook(PBI_XLSX)
    tabled = _tables(wb)
    csvs = {p.stem for p in PBI_DATA.glob("*.csv")}
    assert set(tabled) == csvs, "every CSV should have a matching table, and vice versa"
    for ws in wb.worksheets:
        assert len(ws.tables) <= 1, f"{ws.title} holds more than one table"


def test_powerbi_sheet_names_never_collide_with_table_names(pbi_built):
    """
    Excel allows it; Power Query does not cope with it. When a sheet and a table
    share a name the Navigator appends the table's id - "dim_year" arrives as
    "dim_year1" - and every DAX measure written against the clean name breaks.
    """
    from openpyxl import load_workbook
    wb = load_workbook(PBI_XLSX)
    sheets = set(wb.sheetnames)
    clashes = sorted(sheets & set(_tables(wb)))
    assert not clashes, f"sheet names collide with table names: {clashes}"


def test_powerbi_readme_sheet_is_not_a_table(pbi_built):
    """It must show under Sheets, not Tables, so it is easy to leave unticked."""
    from openpyxl import load_workbook
    wb = load_workbook(PBI_XLSX)
    assert "_README" in wb.sheetnames
    assert not wb["_README"].tables


def test_powerbi_workbook_matches_the_csvs_cell_for_cell(pbi_built):
    from openpyxl import load_workbook
    wb = load_workbook(PBI_XLSX)
    for table, ws in _tables(wb).items():
        rows = list(_csv.reader(open(PBI_DATA / f"{table}.csv")))
        xl = [[c.value for c in r] for r in ws.iter_rows()]
        assert xl[0] == rows[0], f"{table}: header mismatch"
        assert len(xl) == len(rows), f"{table}: row count mismatch"
        for i in range(1, len(rows)):
            for j in range(len(rows[0])):
                a, b = xl[i][j], rows[i][j]
                if a is None:
                    assert b == "", f"{table} r{i}c{j}: blank in xlsx, {b!r} in csv"
                elif isinstance(a, (int, float)):
                    assert _math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-9), \
                        f"{table} r{i}c{j}: {a} vs {b}"
                else:
                    assert str(a) == b, f"{table} r{i}c{j}: {a!r} vs {b!r}"


def test_powerbi_numeric_columns_are_not_text(pbi_built):
    """
    A numeric column containing an empty string is typed as text by Power Query,
    which silently breaks every measure built on it. Blanks must be blank.
    """
    from openpyxl import load_workbook
    wb = load_workbook(PBI_XLSX)
    ws = _tables(wb)["fact_sensitivity"]   # the one table with genuine missing values
    col = [c.value for c in ws["M"]][1:]  # YearsToReserveFloor
    assert any(v is None for v in col), "expected some blanks in this column"
    assert all(v is None or isinstance(v, (int, float)) for v in col), \
        "missing values were written as empty strings rather than blanks"


# ==========================================================================
# 13. The DAX library can actually be created
#
# A measure library is only a text file until Power BI accepts it, and the ways
# it gets refused are all invisible to a reader. A measure may not share its
# name with a column of its home table; names must be unique across the model;
# a "VAR x =" sitting in column 0 looks exactly like the start of a new measure
# to anything parsing the file by eye or by regex; and time-intelligence
# functions need a marked date table, which this model does not have because
# its year key is a whole number.
#
# Every one of those was present in the library at some point and none of them
# showed up as a wrong number - they showed up as a create that failed.
# ==========================================================================

PBI_DAX = ROOT / "powerbi" / "measures.dax"
PBI_DAX_QUERY = ROOT / "powerbi" / "measures_dax_query.dax"


def _measure_defs():
    """(home table, name) for every measure in the query-view file."""
    return _re.findall(r"MEASURE '([^']+)'\[([^\]]+)\]", PBI_DAX_QUERY.read_text())


def test_both_dax_files_hold_the_same_library(pbi_built):
    import build_08_powerbi as B
    paste = {n for n, _ in B._parse_measures(PBI_DAX.read_text())}
    query = {n for _, n in _measure_defs()}
    assert paste == query, (
        "the two forms of the library disagree: "
        f"only in measures.dax {sorted(paste - query)}, "
        f"only in the query form {sorted(query - paste)}"
    )
    assert len(paste) > 90


def test_measure_names_are_unique(pbi_built):
    names = [n for _, n in _measure_defs()]
    dupes = sorted({n for n in names if names.count(n) > 1})
    assert not dupes, f"a model cannot hold two measures with one name: {dupes}"


def test_no_measure_shares_a_name_with_a_column_of_its_home_table(pbi_built):
    """
    Tabular forbids it, and the create fails outright. Nine measures here
    aggregate a column of the same name, which is why they carry a Total prefix.
    """
    cols = {p.stem: next(_csv.reader(open(p))) for p in PBI_DATA.glob("*.csv")}
    clashes = [(t, n) for t, n in _measure_defs() if n in cols.get(t, [])]
    assert not clashes, f"measure name collides with a column of its table: {clashes}"


def test_no_measure_is_named_after_a_dax_keyword(pbi_built):
    """
    "VAR Prior =" in column 0 is a statement inside a measure, not a new
    measure. A parser that misses the distinction splits the measure in two and
    the remainder becomes a measure called "VAR Prior" - which then appears
    three times, because three measures use that variable name.
    """
    import build_08_powerbi as B
    bad = [n for _, n in _measure_defs() if n.startswith(B.DAX_KEYWORDS)]
    assert not bad, f"a DAX statement was parsed as a measure name: {bad}"


def test_every_measure_in_the_query_form_has_a_body(pbi_built):
    text = PBI_DAX_QUERY.read_text()
    empty = [
        n for t, n in _measure_defs()
        if not _re.search(rf"MEASURE '{t}'\[{_re.escape(n)}\] =\n {{8}}\S", text)
    ]
    assert not empty, f"measures written with no expression: {empty}"


def test_every_measure_reference_resolves(pbi_built):
    """
    [Something] in a body must be a measure defined here, a column of some
    table, or an extension column introduced by ADDCOLUMNS in the same measure.
    Anything else is a typo that Power BI reports one measure at a time.
    """
    text = PBI_DAX_QUERY.read_text()
    names = {n for _, n in _measure_defs()}
    cols = {c for p in PBI_DATA.glob("*.csv") for c in next(_csv.reader(open(p)))}
    extension = set(_re.findall(r'"(@[A-Za-z0-9_]+)"', text))
    bodies = _re.sub(r"MEASURE '[^']+'\[[^\]]+\] =", "", text)
    refs = set(_re.findall(r"(?<![A-Za-z0-9_'])\[([^\]]+)\]", bodies))
    unknown = sorted(refs - names - cols - extension)
    assert not unknown, f"references that resolve to nothing: {unknown}"


def test_no_time_intelligence_without_a_date_table(pbi_built):
    """
    dim_year[Year] is a whole number, and the model marks no date table, so
    DATEADD and its relatives error on create. Prior-year comparisons offset
    the integer key instead.
    """
    forbidden = ("DATEADD", "SAMEPERIODLASTYEAR", "PARALLELPERIOD", "DATESYTD",
                 "TOTALYTD", "PREVIOUSYEAR", "DATESBETWEEN")
    for path in (PBI_DAX, PBI_DAX_QUERY):
        text = path.read_text()
        used = [f for f in forbidden if f in text]
        assert not used, f"{path.name} uses {used} but no date table is marked"


def test_the_build_guide_quotes_the_real_measure_count(pbi_built):
    """
    The guide used to claim 103 because the counting regex also matched the
    VAR lines. A number in a document nobody recomputes drifts silently.
    """
    import build_08_powerbi as B
    n = len(B._parse_measures(PBI_DAX.read_text()))
    guide = (ROOT / "powerbi" / "BUILD_GUIDE.md").read_text()
    assert f"{n} measures in twelve groups" in guide, \
        f"the guide does not quote the actual count of {n}"
    assert "103" not in guide, "a stale measure count survives in the guide"


def test_no_variable_is_named_after_a_reserved_word(pbi_built):
    """
    The DAX query parser reserves words the expression parser does not, so a
    measure that Power BI Desktop accepts can still be rejected inside a DEFINE
    block. "VAR Top" was the one that bit: valid as a measure, refused as a
    query with "The syntax for 'Top' is incorrect".
    """
    reserved = {
        "top", "total", "order", "by", "start", "at", "define", "measure",
        "column", "table", "evaluate", "var", "return", "asc", "desc", "skip",
        "dense", "not", "in", "true", "false", "blank", "row", "sample",
    }
    for path in (PBI_DAX, PBI_DAX_QUERY):
        names = _re.findall(r"\bVAR\s+([A-Za-z_][A-Za-z0-9_]*)", path.read_text())
        clashes = sorted({n for n in names if n.lower() in reserved})
        assert not clashes, f"{path.name}: variable named after a reserved word: {clashes}"
